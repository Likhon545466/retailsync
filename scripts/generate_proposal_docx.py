#!/usr/bin/env python3
"""
Full Academic Capstone Project Proposal Document Builder for SHWMS.
Department of Software Engineering, Daffodil International University (DIU).
Course: SE-231 (Software System Analysis & Design / Capstone Project 2).
Part 1: Project Planning and Definition.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# --- Color Palette ---
COLOR_PRIMARY_NAVY = RGBColor(0x1B, 0x36, 0x5D)    # #1B365D - Deep Navy
COLOR_SECONDARY_SLATE = RGBColor(0x2B, 0x4C, 0x7E) # #2B4C7E - Slate Blue
COLOR_DARK_TEXT = RGBColor(0x22, 0x22, 0x22)       # #222222 - Charcoal Text
COLOR_MUTED_GREY = RGBColor(0x55, 0x55, 0x55)      # #555555 - Subtitles & notes
COLOR_HIGHLIGHT_TEAL = RGBColor(0x0D, 0x5C, 0x75)  # #0D5C75 - Accents

HEX_NAVY = "1B365D"
HEX_LIGHT_BLUE = "F0F4F8"
HEX_ALT_ROW = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BORDER = "1B365D"
HEX_CALLOUT_BG = "F1F5F9"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_callout(doc, text_list, title=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_CALLOUT_BORDER}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(10)
        run_title.font.bold = True
        run_title.font.color.rgb = COLOR_PRIMARY_NAVY
        
    for i, t in enumerate(text_list):
        if i > 0 or title:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
        run = p.add_run(t)
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.italic = True
        run.font.color.rgb = COLOR_DARK_TEXT
    
    return tbl

def setup_page_layout(section):
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    
    section.different_first_page_header_footer = True
    
    # Running Header (pages 2+)
    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(2)
    hrun = hp.add_run("RetailSync: Centralized Super Shop WMS — Project Proposal (DIU SE-231)")
    hrun.font.name = "Calibri"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = COLOR_MUTED_GREY
    
    # Running Footer (pages 2+)
    from docx.enum.text import WD_TAB_ALIGNMENT
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.paragraph_format.space_before = Pt(4)
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.8), WD_TAB_ALIGNMENT.RIGHT)
    
    frun_left = fp.add_run("Department of Software Engineering, Daffodil International University")
    frun_left.font.name = "Calibri"
    frun_left.font.size = Pt(8.5)
    frun_left.font.color.rgb = COLOR_MUTED_GREY
    
    frun_tab = fp.add_run("\tPage ")
    frun_tab.font.name = "Calibri"
    frun_tab.font.size = Pt(8.5)
    frun_tab.font.color.rgb = COLOR_MUTED_GREY
    
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp._p.append(fldSimple)

def add_part_header(doc, part_letter, part_title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run_part = p.add_run(f"PART {part_letter}\n")
    run_part.font.name = "Calibri"
    run_part.font.size = Pt(11)
    run_part.font.bold = True
    run_part.font.color.rgb = COLOR_HIGHLIGHT_TEAL
    
    run_title = p.add_run(part_title.upper())
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(15)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY_NAVY

def add_h1(doc, title, page_break_before=False):
    p = doc.add_paragraph()
    if page_break_before:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY_NAVY
    return p

def add_h2(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY_SLATE
    return p

def add_h3(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_p(doc, text="", bold_prefix="", space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK_TEXT
    if text:
        r_text = p.add_run(text)
        r_text.font.name = "Calibri"
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_bullet(doc, text="", bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK_TEXT
    if text:
        r_text = p.add_run(text)
        r_text.font.name = "Calibri"
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = COLOR_DARK_TEXT
    return p

def build_styled_table(doc, headers, data, col_widths=None, alignment=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, HEX_BORDER)
    
    # Header Row
    hdr_row = tbl.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for i, title in enumerate(headers):
        cell = hdr_row.cells[i]
        set_cell_background(cell, HEX_NAVY)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if alignment and i < len(alignment):
            p.alignment = alignment[i]
        run = p.add_run(title)
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    # Data Rows
    for r_idx, row_data in enumerate(data):
        row = tbl.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        bg_color = HEX_ALT_ROW if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if alignment and c_idx < len(alignment):
                p.alignment = alignment[c_idx]
            run = p.add_run(str(val))
            run.font.name = "Calibri"
            run.font.size = Pt(9.0)
            run.font.color.rgb = COLOR_DARK_TEXT
            
    # Set Column Widths if provided
    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
                
    return tbl

print("Table and text formatting setup complete.")
