#!/usr/bin/env python3
"""
RetailSync Engineering Documentation Markdown to DOCX Converter.
Converts all Markdown specifications into university-standard Microsoft Word documents.
Course: SE-231 (Software System Analysis & Design / Capstone Project 2)
Author: Raisul Islam Likhon (Section: SWE-44D)
Department of Software Engineering, Daffodil International University (DIU)
"""

import os
import re
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# --- Color Palette (DIU Academic Design System) ---
COLOR_PRIMARY_NAVY = RGBColor(0x1B, 0x36, 0x5D)    # #1B365D - Deep Navy
COLOR_SECONDARY_SLATE = RGBColor(0x2B, 0x4C, 0x7E) # #2B4C7E - Slate Blue
COLOR_DARK_TEXT = RGBColor(0x22, 0x22, 0x22)       # #222222 - Charcoal Text
COLOR_MUTED_GREY = RGBColor(0x55, 0x55, 0x55)      # #555555 - Subtitles & notes
COLOR_HIGHLIGHT_TEAL = RGBColor(0x0D, 0x5C, 0x75)  # #0D5C75 - Accents
COLOR_CODE_DARK = RGBColor(0x1E, 0x29, 0x3B)       # #1E293B - Code text

HEX_NAVY = "1B365D"
HEX_LIGHT_BLUE = "F0F4F8"
HEX_ALT_ROW = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BORDER = "1B365D"
HEX_CALLOUT_BG = "F1F5F9"
HEX_CODE_BG = "F8FAFC"
HEX_CODE_BORDER = "E2E8F0"

PAGE_WIDTH_INCHES = 6.5  # 8.5 total minus 2x 1.0 inch margins

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def setup_page_layout(section, header_title="RetailSync Engineering Specifications"):
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.different_first_page_header_footer = True

    # Header
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(4)
    hrun = hp.add_run(header_title)
    hrun.font.name = "Calibri"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = COLOR_MUTED_GREY

    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.paragraph_format.space_before = Pt(4)
    frun = fp.add_run("Dept. of Software Engineering, Daffodil International University | Raisul Islam Likhon")
    frun.font.name = "Calibri"
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = COLOR_MUTED_GREY

    frun_tab = fp.add_run("\t\tPage ")
    frun_tab.font.name = "Calibri"
    frun_tab.font.size = Pt(8.5)
    frun_tab.font.color.rgb = COLOR_MUTED_GREY

    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp._p.append(fldSimple)

def clean_text_symbols(text):
    """Replaces raw LaTeX math notation with clean Unicode symbols for Word."""
    if not text:
        return text
    text = re.sub(r'\$\\le\s*', '≤ ', text)
    text = re.sub(r'\\le\s*', '≤ ', text)
    text = re.sub(r'\$\\ge\s*', '≥ ', text)
    text = re.sub(r'\\ge\s*', '≥ ', text)
    text = re.sub(r'\$<\s*', '< ', text)
    text = re.sub(r'\$>\s*', '> ', text)
    text = re.sub(r'\\%', '%', text)
    text = re.sub(r'\$', '', text)
    return text

def parse_inline(p, text, base_font="Calibri", base_size=Pt(10), base_color=COLOR_DARK_TEXT, default_bold=False):
    """
    Parses markdown inline styling: **bold**, *italic*, and `code`.
    Adds runs directly to paragraph p.
    """
    if not text:
        return
    
    text = clean_text_symbols(text)
    
    # Regex pattern to match **bold**, *italic*, `code`
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    parts = pattern.split(text)

    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            clean = part[2:-2]
            run = p.add_run(clean)
            run.font.name = base_font
            run.font.size = base_size
            run.font.bold = True
            run.font.color.rgb = base_color
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2 and not part.startswith('**'):
            clean = part[1:-1]
            run = p.add_run(clean)
            run.font.name = base_font
            run.font.size = base_size
            run.font.italic = True
            run.font.color.rgb = base_color
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            clean = part[1:-1]
            run = p.add_run(clean)
            run.font.name = "Consolas"
            run.font.size = Pt(base_size.pt - 0.5 if hasattr(base_size, 'pt') else 9)
            run.font.color.rgb = COLOR_HIGHLIGHT_TEAL
        else:
            run = p.add_run(part)
            run.font.name = base_font
            run.font.size = base_size
            if default_bold:
                run.font.bold = True
            run.font.color.rgb = base_color

def add_code_block(doc, code_lines, language=""):
    """Creates a beautifully shaded code/diagram callout block."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_CODE_BG)
    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="12" w:space="0" w:color="{HEX_NAVY}"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_CODE_BORDER}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_CODE_BORDER}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_CODE_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05

    code_text = "\n".join(code_lines)
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = COLOR_CODE_DARK

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_callout(doc, text_lines, title=""):
    """Creates a shaded blockquote callout with left accent border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=100, bottom=100, left=180, right=140)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_CALLOUT_BORDER}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15

    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(10.5)
        run_title.font.bold = True
        run_title.font.color.rgb = COLOR_PRIMARY_NAVY

    for i, t in enumerate(text_lines):
        if i > 0 or title:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
        parse_inline(p, t, base_font="Calibri", base_size=Pt(9.5), base_color=COLOR_DARK_TEXT)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_markdown_table(doc, table_rows):
    """
    Renders a parsed markdown table with styled headers, alternating rows,
    clean cell borders, and intelligent column widths.
    """
    if not table_rows or len(table_rows) < 2:
        return

    # First row is header
    header_cols = [c.strip() for c in table_rows[0]]
    # Check if second row is separator row e.g. :--- | :---
    data_rows = []
    start_idx = 1
    if len(table_rows) > 1 and all(set(c.strip()).issubset({'-', ':', ' '}) for c in table_rows[1]):
        start_idx = 2

    for r in table_rows[start_idx:]:
        data_rows.append([c.strip() for c in r])

    num_cols = len(header_cols)
    tbl = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, HEX_BORDER)

    # Header Row
    hdr_row = tbl.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for c_idx, title in enumerate(header_cols):
        cell = hdr_row.cells[c_idx]
        set_cell_background(cell, HEX_NAVY)
        set_cell_margins(cell, top=100, bottom=100, left=110, right=110)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        parse_inline(p, title, base_font="Calibri", base_size=Pt(9.0), base_color=RGBColor(0xFF, 0xFF, 0xFF), default_bold=True)

    # Data Rows
    for r_idx, row_vals in enumerate(data_rows):
        row = tbl.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg_color = HEX_ALT_ROW if (r_idx % 2 == 1) else "FFFFFF"

        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            val = row_vals[c_idx] if c_idx < len(row_vals) else ""
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=65, bottom=65, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            parse_inline(p, val, base_font="Calibri", base_size=Pt(8.5), base_color=COLOR_DARK_TEXT)

    # Calculate proportional column widths
    max_lens = [max(len(header_cols[i]), 1) for i in range(num_cols)]
    for row_vals in data_rows:
        for i in range(num_cols):
            if i < len(row_vals):
                max_lens[i] = max(max_lens[i], len(row_vals[i]))

    # Adjust weights: constrain very long columns and ensure minimum width
    weights = [min(max(l, 8), 50) for l in max_lens]
    total_weight = sum(weights)
    col_widths = [(w / total_weight) * PAGE_WIDTH_INCHES for w in weights]

    for row in tbl.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

def convert_markdown_to_docx(md_path, docx_path, doc_title=None, is_subdoc=False, parent_doc=None):
    """
    Parses a markdown document and writes it into a styled Word Document.
    Can be used standalone or appended into a parent combined document.
    """
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    if parent_doc is not None:
        doc = parent_doc
    else:
        doc = docx.Document()
        section = doc.sections[0]
        header_text = doc_title or "RetailSync Engineering Specification"
        setup_page_layout(section, header_title=header_text)

    in_code_block = False
    code_block_lang = ""
    code_block_lines = []

    in_table = False
    table_rows = []

    in_callout = False
    callout_lines = []

    idx = 0
    total_lines = len(lines)

    while idx < total_lines:
        line = lines[idx].rstrip('\r\n')
        stripped = line.strip()

        # Handle Code Fences
        if stripped.startswith('```'):
            if in_code_block:
                add_code_block(doc, code_block_lines, code_block_lang)
                in_code_block = False
                code_block_lines = []
                code_block_lang = ""
            else:
                # Flush table if any
                if in_table:
                    add_markdown_table(doc, table_rows)
                    in_table = False
                    table_rows = []
                # Flush callout if any
                if in_callout:
                    add_callout(doc, callout_lines)
                    in_callout = False
                    callout_lines = []

                in_code_block = True
                code_block_lang = stripped[3:].strip()
                code_block_lines = []
            idx += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            idx += 1
            continue

        # Handle Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            # Parse row columns
            cols = [c.strip() for c in stripped.strip('|').split('|')]
            if not in_table:
                # Flush callout
                if in_callout:
                    add_callout(doc, callout_lines)
                    in_callout = False
                    callout_lines = []
                in_table = True
                table_rows = [cols]
            else:
                table_rows.append(cols)
            idx += 1
            continue
        else:
            if in_table:
                add_markdown_table(doc, table_rows)
                in_table = False
                table_rows = []

        # Handle Blockquotes
        if stripped.startswith('>'):
            quote_text = stripped[1:].strip()
            if not in_callout:
                in_callout = True
                callout_lines = [quote_text]
            else:
                callout_lines.append(quote_text)
            idx += 1
            continue
        else:
            if in_callout:
                add_callout(doc, callout_lines)
                in_callout = False
                callout_lines = []

        # Blank line
        if not stripped:
            idx += 1
            continue

        # Horizontal Rule
        if stripped in ('---', '***', '___'):
            # Render subtle divider or minor spacing
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="{HEX_BORDER}"/></w:pBdr>')
            pPr.append(pBdr)
            idx += 1
            continue

        # Headings
        if stripped.startswith('# ') and not stripped.startswith('## '):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(title_text)
            run.font.name = "Calibri"
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY_NAVY
            idx += 1
            continue

        if stripped.startswith('## '):
            h2_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h2_text)
            run.font.name = "Calibri"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY_NAVY
            idx += 1
            continue

        if stripped.startswith('### '):
            h3_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h3_text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = COLOR_SECONDARY_SLATE
            idx += 1
            continue

        if stripped.startswith('#### '):
            h4_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h4_text)
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = COLOR_DARK_TEXT
            idx += 1
            continue

        # Bullet List Items (*, -, +)
        bullet_match = re.match(r'^(\s*)([\*\-\+])\s+(.*)$', line)
        if bullet_match:
            indent_spaces = len(bullet_match.group(1))
            bullet_content = bullet_match.group(3)
            
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.15
            if indent_spaces >= 2:
                p.paragraph_format.left_indent = Inches(0.4)
            parse_inline(p, bullet_content, base_font="Calibri", base_size=Pt(9.5), base_color=COLOR_DARK_TEXT)
            idx += 1
            continue

        # Numbered List Items (1. , 2. )
        num_match = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
        if num_match:
            num_str = num_match.group(2)
            num_content = num_match.group(3)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.15
            parse_inline(p, num_content, base_font="Calibri", base_size=Pt(9.5), base_color=COLOR_DARK_TEXT)
            idx += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        parse_inline(p, stripped, base_font="Calibri", base_size=Pt(10), base_color=COLOR_DARK_TEXT)
        idx += 1

    # Cleanup remaining state
    if in_code_block:
        add_code_block(doc, code_block_lines, code_block_lang)
    if in_table:
        add_markdown_table(doc, table_rows)
    if in_callout:
        add_callout(doc, callout_lines)

    if parent_doc is None:
        doc.save(docx_path)
        print(f"Successfully generated: {docx_path}")

    return doc

def build_master_engineering_specification():
    """
    Builds a unified, comprehensive engineering specification DOCX that merges
    all 6 documents into a single master engineering blueprint for DIU Capstone.
    """
    print("\n--- Compiling Unified Master Engineering Specification Document ---")
    doc = docx.Document()
    section = doc.sections[0]
    setup_page_layout(section, header_title="RetailSync: Complete Engineering Specifications | DIU SWE Capstone")

    # Cover Page
    cp = doc.add_paragraph()
    cp.paragraph_format.space_before = Pt(10)
    cp.paragraph_format.space_after = Pt(2)
    run_dept = cp.add_run("DEPARTMENT OF SOFTWARE ENGINEERING\nDAFFODIL INTERNATIONAL UNIVERSITY")
    run_dept.font.name = "Calibri"
    run_dept.font.size = Pt(10.5)
    run_dept.font.bold = True
    run_dept.font.color.rgb = COLOR_HIGHLIGHT_TEAL

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("RetailSync: Super Shop Warehouse Management System")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(21)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY_NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(10)
    run_sub = p_sub.add_run("Comprehensive Engineering Specification Suite & Technical Blueprint")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_SECONDARY_SLATE

    # Metadata Callout Box
    meta_lines = [
        "Course: SE-231 (Software System Analysis & Design / Capstone Project 2)",
        "Student / Author: Raisul Islam Likhon (Section: SWE-44D)",
        "Institution: Daffodil International University (DIU), Dhaka, Bangladesh",
        "Submission Term: Fall 2026 | Version: 1.0.0-RELEASE (Master Production Specification)"
    ]
    add_callout(doc, meta_lines, title="ACADEMIC PROJECT METADATA")

    # Document Modules Table of Contents
    p_toc_head = doc.add_paragraph()
    p_toc_head.paragraph_format.space_before = Pt(10)
    p_toc_head.paragraph_format.space_after = Pt(4)
    r = p_toc_head.add_run("Specification Modules Included in this Volume")
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY_NAVY

    modules_matrix = [
        ["Module #", "Document Name", "Primary Technical Focus", "Target Output Format"],
        ["Module 1", "Product Requirements Document (PRD)", "Personas, User Journeys, Epics, Quantitative SLOs", "PRD.md / PRD.docx"],
        ["Module 2", "Software Architecture Document (SAD)", "4-Tier Architecture, Concurrency, FEFO Data Flows", "ARCHITECTURE.md / ARCHITECTURE.docx"],
        ["Module 3", "Technology Stack & ADRs", "Next.js 14, FastAPI, PostgreSQL 16, Redis 7, Hardware", "TECH_STACK.md / TECH_STACK.docx"],
        ["Module 4", "Database Schema & Data Dictionary", "3NF Relational Model, Triggers, Indexes, DDL Tables", "DATABASE_SCHEMA.md / DATABASE_SCHEMA.docx"],
        ["Module 5", "REST API Specification & Contract", "OpenAPI 3.1 Endpoints, Request/Response Envelopes", "API_SPECIFICATION.md / API_SPECIFICATION.docx"],
        ["Module 6", "Sprint Implementation Plan & Tasks", "14-Week Agile Roadmap, Story Points, DoD Criteria", "SPRINT_PLAN_AND_TASKS.md / SPRINT_PLAN_AND_TASKS.docx"]
    ]
    add_markdown_table(doc, modules_matrix)

    # Determine absolute project root
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(script_dir)

    # Append each subdocument
    spec_files = [
        (os.path.join(project_root, "docs/01_Product_Requirements/PRD.md"), "Module 1: Product Requirements Document (PRD)"),
        (os.path.join(project_root, "docs/02_Software_Architecture/ARCHITECTURE.md"), "Module 2: Software Architecture Document (SAD)"),
        (os.path.join(project_root, "docs/03_Technology_Stack/TECH_STACK.md"), "Module 3: Technology Stack & Architectural Decision Records"),
        (os.path.join(project_root, "docs/04_Database_Design/DATABASE_SCHEMA.md"), "Module 4: Relational Database Schema & Data Dictionary"),
        (os.path.join(project_root, "docs/05_API_Specifications/API_SPECIFICATION.md"), "Module 5: REST API Specification & Endpoint Contracts"),
        (os.path.join(project_root, "docs/06_Sprint_Planning/SPRINT_PLAN_AND_TASKS.md"), "Module 6: Sprint Implementation Plan & Task Breakdown")
    ]

    for file_path, module_title in spec_files:
        if not os.path.exists(file_path):
            print(f"Warning: Module file {file_path} not found.")
            continue
        doc.add_page_break()
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(18)
        p_sec.paragraph_format.space_after = Pt(6)
        r_sec = p_sec.add_run(module_title.upper())
        r_sec.font.name = "Calibri"
        r_sec.font.size = Pt(16)
        r_sec.font.bold = True
        r_sec.font.color.rgb = COLOR_PRIMARY_NAVY

        convert_markdown_to_docx(file_path, None, is_subdoc=True, parent_doc=doc)

    # Add DDL Appendix to Master Document
    sql_path = os.path.join(project_root, "docs/04_Database_Design/DATABASE_SCHEMA.sql")
    if os.path.exists(sql_path):
        doc.add_page_break()
        p_sql_title = doc.add_paragraph()
        p_sql_title.paragraph_format.space_before = Pt(18)
        p_sql_title.paragraph_format.space_after = Pt(6)
        r_sql = p_sql_title.add_run("APPENDIX A: COMPLETE PRODUCTION POSTGRESQL 16 DDL SCRIPT")
        r_sql.font.name = "Calibri"
        r_sql.font.size = Pt(15)
        r_sql.font.bold = True
        r_sql.font.color.rgb = COLOR_PRIMARY_NAVY

        with open(sql_path, 'r', encoding='utf-8') as f:
            sql_lines = [line.rstrip('\r\n') for line in f.readlines()]
        
        add_code_block(doc, sql_lines, language="sql")

    master_path = os.path.join(project_root, "docs/RetailSync_Master_Engineering_Suite.docx")
    doc.save(master_path)
    print(f"Master Document Successfully Saved: {master_path}")

def main():
    print("=================================================================")
    print("RetailSync Markdown to Word (.docx) Documentation Suite Converter")
    print("=================================================================")

    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(script_dir)

    docs_to_convert = [
        (os.path.join(project_root, "docs/01_Product_Requirements/PRD.md"),
         os.path.join(project_root, "docs/01_Product_Requirements/PRD.docx"),
         "RetailSync: Product Requirements Document (PRD)"),

        (os.path.join(project_root, "docs/02_Software_Architecture/ARCHITECTURE.md"),
         os.path.join(project_root, "docs/02_Software_Architecture/ARCHITECTURE.docx"),
         "RetailSync: Software Architecture Document (SAD)"),

        (os.path.join(project_root, "docs/03_Technology_Stack/TECH_STACK.md"),
         os.path.join(project_root, "docs/03_Technology_Stack/TECH_STACK.docx"),
         "RetailSync: Technology Stack & ADRs"),

        (os.path.join(project_root, "docs/04_Database_Design/DATABASE_SCHEMA.md"),
         os.path.join(project_root, "docs/04_Database_Design/DATABASE_SCHEMA.docx"),
         "RetailSync: Relational Database Schema & Data Dictionary"),

        (os.path.join(project_root, "docs/05_API_Specifications/API_SPECIFICATION.md"),
         os.path.join(project_root, "docs/05_API_Specifications/API_SPECIFICATION.docx"),
         "RetailSync: REST API Specification & Endpoint Contracts"),

        (os.path.join(project_root, "docs/06_Sprint_Planning/SPRINT_PLAN_AND_TASKS.md"),
         os.path.join(project_root, "docs/06_Sprint_Planning/SPRINT_PLAN_AND_TASKS.docx"),
         "RetailSync: Sprint Implementation Plan & Task Breakdown"),
    ]

    for md_file, docx_file, title in docs_to_convert:
        if os.path.exists(md_file):
            convert_markdown_to_docx(md_file, docx_file, doc_title=title)
        else:
            print(f"Warning: File {md_file} not found.")

    # Also generate the unified master engineering specification document
    build_master_engineering_specification()
    print("\nAll DOCX documents generated successfully!")

if __name__ == '__main__':
    main()
